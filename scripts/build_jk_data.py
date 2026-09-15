import openpyxl
import json
import re
import os

path = r'C:\Users\pavan\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\F9C8E35F817C89D982A1A2C62B262666B2E51205\transfers\2026-37\IKS_Traditional_Measurements_JammuKashmir_Expanded.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

sheet_mapping = {
    'Transportation & Distance': ('transportation-distance', 'trans'),
    'Land Measurement': ('land-measurement', 'land'),
    'Livestock & Dairy': ('livestock-dairy', 'dairy'),
    'Household & Daily Life': ('household', 'hh'),
    'Gold & Jewellery': ('gold-jewellery', 'gold'),
    'Seed & Crop (Agriculture)': ('agriculture', 'agri'),
    'Currency & Money': ('currency-money', 'curr'),
    'Storage & Transportation': ('storage-transport', 'storage'),
    'Religious & Cultural Sectors': ('religious-cultural', 'relig'),
    'Trade & Commerce': ('trade-commerce', 'trade'),
    'Textile & Handloom': ('textile-handloom', 'textile'),
    'Medicine (Ayurveda)': ('medicine', 'med'),
    'Construction & Architecture': ('architecture', 'arch'),
    'Time & Calendar': ('time', 'time'),
}

def clean_slug(text):
    s = re.sub(r'[^a-zA-Z0-9]+', '-', str(text).lower()).strip('-')
    return s if s else "unit"

def map_category(cat_str):
    if not cat_str or cat_str == "N/A": return "other"
    c = str(cat_str).lower()
    if "weight" in c: return "weight"
    if "volume" in c or "capacity" in c: return "volume"
    if "length" in c or "distance" in c: return "length"
    if "area" in c or "land" in c: return "area"
    if "currency" in c or "exchange" in c or "coin" in c or "money" in c: return "currency"
    if "time" in c or "calendar" in c: return "time"
    return "other"

all_measurements = []

for sheet_name, (sector_slug, sector_code) in sheet_mapping.items():
    ws = wb[sheet_name]
    sheet_units = []
    
    for r in range(3, ws.max_row + 1):
        unit_name_val = ws.cell(r, 2).value
        if unit_name_val is None or not str(unit_name_val).strip():
            continue
        
        unit_name = str(unit_name_val).strip()
        sanskrit_name = str(ws.cell(r, 3).value).strip() if ws.cell(r, 3).value is not None else ""
        local_name = str(ws.cell(r, 4).value).strip() if ws.cell(r, 4).value is not None else ""
        hindi_name = str(ws.cell(r, 5).value).strip() if ws.cell(r, 5).value is not None else ""
        meas_type = str(ws.cell(r, 6).value).strip() if ws.cell(r, 6).value is not None else ""
        approx_eq = str(ws.cell(r, 7).value).strip() if ws.cell(r, 7).value is not None else ""
        relation = str(ws.cell(r, 8).value).strip() if ws.cell(r, 8).value is not None else ""
        used_in_val = str(ws.cell(r, 9).value).strip() if ws.cell(r, 9).value is not None else ""
        ref_url = str(ws.cell(r, 10).value).strip() if ws.cell(r, 10).value is not None else ""
        entry_source = str(ws.cell(r, 11).value).strip() if ws.cell(r, 11).value is not None else ""

        row_idx = len(sheet_units) + 1
        item_id = f"jk-{sector_code}-{row_idx}"
        slug = f"jk-{sector_code}-{clean_slug(unit_name)}-{row_idx}"
        category = map_category(meas_type)

        item = {
            "id": item_id,
            "slug": slug,
            "name_english": unit_name,
            "category": category,
            "sector": sector_slug,
            "origin": "Jammu & Kashmir",
            "states": ["Jammu and Kashmir", "Jammu & Kashmir"],
            "created_at": "2024-01-01"
        }

        # Sanskrit Name
        if sanskrit_name and sanskrit_name != "None" and sanskrit_name != "N/A":
            item["name_sanskrit"] = sanskrit_name

        # Local Language Name (Kashmiri/Dogri)
        if local_name and local_name != "None" and local_name != "—" and local_name != "N/A":
            item["local_names"] = [local_name]

        # Hindi Name
        if hindi_name and hindi_name != "None" and hindi_name != "—" and hindi_name != "N/A":
            item["name_hindi"] = hindi_name

        # Type / Measurement
        if meas_type and meas_type != "None" and meas_type != "N/A":
            item["measurement_type"] = meas_type

        # Approx. Modern Equivalent
        if approx_eq and approx_eq != "None" and approx_eq != "—" and approx_eq != "N/A":
            item["modern_equivalent"] = approx_eq

        # Relation / Hierarchy
        if relation and relation != "None" and relation != "—" and relation != "N/A":
            item["conversion_formula"] = relation

        # Used In / Context
        if used_in_val and used_in_val != "None" and used_in_val != "—" and used_in_val != "N/A":
            item["meaning"] = used_in_val
            item["historical_context"] = used_in_val
            item["used_in"] = [used_in_val]

        # References
        if ref_url and ref_url != "None" and ref_url != "—" and ref_url != "N/A":
            item["references"] = [ref_url]

        item["tags"] = [
            "jammu-and-kashmir",
            "kashmir",
            "dogri",
            "traditional-units",
            sector_slug,
            clean_slug(unit_name),
            category
        ]

        sheet_units.append(item)
    
    all_measurements.extend(sheet_units)
    print(f"Sector '{sector_slug}' ({sheet_name}): {len(sheet_units)} units")

print(f"\nTotal Jammu & Kashmir Measurements: {len(all_measurements)}")

# Check uniqueness of IDs and Slugs
ids = [m['id'] for m in all_measurements]
assert len(ids) == len(set(ids)), "Duplicate IDs found!"
slugs = [m['slug'] for m in all_measurements]
assert len(slugs) == len(set(slugs)), "Duplicate Slugs found!"
print("All IDs and Slugs are completely unique!")

# Generate TS code
ts_code = '''import { Measurement } from "@/types";

export const JK_MEASUREMENTS: Measurement[] = ''' + json.dumps(all_measurements, ensure_ascii=False, indent=2) + ';\n'

with open('lib/jkData.ts', 'w', encoding='utf-8') as f:
    f.write(ts_code)

print("Successfully written to lib/jkData.ts!")
