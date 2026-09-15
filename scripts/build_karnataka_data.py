import openpyxl
import json
import re
import os

path = r'C:\Users\pavan\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\F9C8E35F817C89D982A1A2C62B262666B2E51205\transfers\2026-37\Karnataka_Traditional_Units.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

sheet_mapping = {
    'Agriculture': ('agriculture', 'agri'),
    'Trade & Commerce': ('trade-commerce', 'trade'),
    'Architecture': ('architecture', 'arch'),
    'Medicine (Ayurveda)': ('medicine', 'med'),
    'Textile & Handloom': ('textile-handloom', 'textile'),
    'Currency & Money': ('currency-money', 'curr'),
    'Household': ('household', 'hh'),
    'Storage & Transportation': ('storage-transport', 'storage'),
    'land measuement': ('land-measurement', 'land'),
    'transportation&distance': ('transportation-distance', 'trans'),
    'livestock&dairy': ('livestock-dairy', 'dairy'),
    'Sheet4': ('gold-jewellery', 'gold'),
}

def clean_slug(text):
    s = re.sub(r'[^a-zA-Z0-9]+', '-', str(text).lower()).strip('-')
    return s if s else "unit"

def map_category(cat_str):
    if not cat_str: return "other"
    c = str(cat_str).lower()
    if "weight" in c: return "weight"
    if "volume" in c or "capacity" in c: return "volume"
    if "length" in c or "distance" in c: return "length"
    if "area" in c: return "area"
    if "currency" in c: return "currency"
    if "time" in c: return "time"
    return "other"

all_measurements = []

for sheet_name, (sector_slug, sector_code) in sheet_mapping.items():
    ws = wb[sheet_name]
    sheet_units = []
    for r in range(2, ws.max_row + 1):
        s_no = ws.cell(r, 1).value
        unit_name = ws.cell(r, 2).value
        if unit_name is None or not str(unit_name).strip():
            continue
        unit_name = str(unit_name).strip()
        kannada_name = str(ws.cell(r, 3).value).strip() if ws.cell(r, 3).value is not None else ""
        english_trans = str(ws.cell(r, 4).value).strip() if ws.cell(r, 4).value is not None else ""
        sanskrit_name = str(ws.cell(r, 5).value).strip() if ws.cell(r, 5).value is not None else ""
        meas_category = str(ws.cell(r, 6).value).strip() if ws.cell(r, 6).value is not None else ""
        order_val = ws.cell(r, 7).value
        relation = str(ws.cell(r, 8).value).strip() if ws.cell(r, 8).value is not None else ""
        modern_si = str(ws.cell(r, 9).value).strip() if ws.cell(r, 9).value is not None else ""
        hist_usage = str(ws.cell(r, 10).value).strip() if ws.cell(r, 10).value is not None else ""
        sector_in_cell = str(ws.cell(r, 11).value).strip() if ws.cell(r, 11).value is not None else ""
        description = str(ws.cell(r, 12).value).strip() if ws.cell(r, 12).value is not None else ""
        hist_period = str(ws.cell(r, 13).value).strip() if ws.cell(r, 13).value is not None else ""
        region_ka = str(ws.cell(r, 14).value).strip() if ws.cell(r, 14).value is not None else "Karnataka"
        source_ref = str(ws.cell(r, 15).value).strip() if ws.cell(r, 15).value is not None else ""

        row_idx = len(sheet_units) + 1
        item_id = f"ka-{sector_code}-{row_idx}"
        slug = f"ka-{sector_code}-{clean_slug(unit_name)}-{row_idx}"
        category = map_category(meas_category)

        item = {
            "id": item_id,
            "slug": slug,
            "name_english": unit_name,
            "category": category,
            "sector": sector_slug,
            "origin": region_ka if region_ka and region_ka != "None" else "Karnataka",
            "states": ["Karnataka"],
            "created_at": "2024-01-01"
        }

        if kannada_name and kannada_name != "None" and kannada_name != "—":
            item["local_names"] = [kannada_name]
        
        if english_trans and english_trans != "None" and english_trans != "—":
            item["name_hindi"] = english_trans # mapped to 4th column in standard table
        
        if sanskrit_name and sanskrit_name != "None" and sanskrit_name != "—":
            item["name_sanskrit"] = sanskrit_name

        if meas_category and meas_category != "None":
            item["measurement_type"] = meas_category

        if modern_si and modern_si != "None":
            item["modern_equivalent"] = modern_si

        if relation and relation != "None":
            item["conversion_formula"] = relation

        if hist_usage and hist_usage != "None":
            item["used_in"] = [hist_usage]

        if description and description != "None":
            item["meaning"] = description
            item["historical_context"] = description

        if hist_period and hist_period != "None":
            item["historical_period"] = hist_period

        if source_ref and source_ref != "None" and source_ref != "—":
            item["references"] = [source_ref]

        item["tags"] = [
            "karnataka",
            "traditional-units",
            sector_slug,
            clean_slug(unit_name),
            category
        ]

        sheet_units.append(item)
    all_measurements.extend(sheet_units)
    print(f"Sector '{sector_slug}' ({sheet_name}): {len(sheet_units)} units")

print(f"\nTotal Karnataka Measurements Generated: {len(all_measurements)}")

# Check ID uniqueness
ids = [m['id'] for m in all_measurements]
assert len(ids) == len(set(ids)), "Duplicate IDs found!"
print("All IDs are unique!")

# Generate TS file
ts_content = '''import { Measurement } from "@/types";

export const KARNATAKA_MEASUREMENTS: Measurement[] = ''' + json.dumps(all_measurements, ensure_ascii=False, indent=2) + ';\n'

with open('lib/karnatakaData.ts', 'w', encoding='utf-8') as f:
    f.write(ts_content)

print("Successfully written to lib/karnatakaData.ts!")
