import openpyxl
import json
import re

path = r'c:\Users\pavan\Downloads\IKS_Traditional_Measurements_Mizoram.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

sheet_mapping = {
    'Transportation & Distance': ('transportation-distance', 'trans'),
    'Land Measurement': ('land-measurement', 'land'),
    'Livestock & Dairy': ('livestock-dairy', 'dairy'),
    'Household & Daily Life': ('household', 'hh'),
    'Gold & Jewellery': ('gold-jewellery', 'gold'),
    'Seed & Crop (Agriculture)': ('agriculture', 'agri'),
    'Currency & Money': ('currency-money', 'curr'),
    'Storage & Transportation': ('storage-transport', 'storage'),
    'Religious & Cultural Sectors': ('religious-cultural', 'relig'),
    'Trade & Commerce': ('trade-commerce', 'trade'),
    'Textile & Handloom': ('textile-handloom', 'textile'),
    'Medicine (Ayurveda)': ('medicine', 'med'),
    'Construction & Architecture': ('architecture', 'arch'),
}

def clean_slug(text):
    s = re.sub(r'[^a-zA-Z0-9]+', '-', str(text).lower()).strip('-')
    return s if s else "unit"

def map_category(cat_str):
    if not cat_str or cat_str == "N/A": return "other"
    c = str(cat_str).lower()
    if "weight" in c: return "weight"
    if "volume" in c or "capacity" in c: return "volume"
    if "length" in c or "distance" in c: return "length"
    if "area" in c: return "area"
    if "currency" in c or "exchange" in c or "wealth" in c: return "currency"
    if "time" in c: return "time"
    return "other"

all_measurements = []

for sheet_name, (sector_slug, sector_code) in sheet_mapping.items():
    ws = wb[sheet_name]
    sheet_units = []
    for r in range(3, ws.max_row + 1):
        num_val = ws.cell(r, 1).value
        unit_name = ws.cell(r, 2).value
        if unit_name is None or not str(unit_name).strip():
            continue
        unit_name = str(unit_name).strip()
        sanskrit_name = str(ws.cell(r, 3).value).strip() if ws.cell(r, 3).value is not None else ""
        local_name = str(ws.cell(r, 4).value).strip() if ws.cell(r, 4).value is not None else ""
        hindi_name = str(ws.cell(r, 5).value).strip() if ws.cell(r, 5).value is not None else ""
        type_meas = str(ws.cell(r, 6).value).strip() if ws.cell(r, 6).value is not None else ""
        approx_eq = str(ws.cell(r, 7).value).strip() if ws.cell(r, 7).value is not None else ""
        relation = str(ws.cell(r, 8).value).strip() if ws.cell(r, 8).value is not None else ""
        used_in = str(ws.cell(r, 9).value).strip() if ws.cell(r, 9).value is not None else ""
        ref_url = str(ws.cell(r, 10).value).strip() if ws.cell(r, 10).value is not None else ""

        row_idx = len(sheet_units) + 1
        item_id = f"mz-{sector_code}-{row_idx}"
        slug = f"mz-{sector_code}-{clean_slug(unit_name)}-{row_idx}"
        category = map_category(type_meas)

        item = {
            "id": item_id,
            "slug": slug,
            "name_english": unit_name,
            "category": category,
            "sector": sector_slug,
            "origin": "Mizoram",
            "states": ["Mizoram"],
            "created_at": "2024-01-01"
        }

        if sanskrit_name and sanskrit_name != "N/A" and sanskrit_name != "None":
            item["name_sanskrit"] = sanskrit_name

        if local_name and local_name != "N/A" and local_name != "None":
            item["local_names"] = [local_name]

        if hindi_name and hindi_name != "N/A" and hindi_name != "None":
            item["name_hindi"] = hindi_name

        if type_meas and type_meas != "N/A" and type_meas != "None":
            item["measurement_type"] = type_meas

        if approx_eq and approx_eq != "N/A" and approx_eq != "None":
            item["modern_equivalent"] = approx_eq

        if relation and relation != "N/A" and relation != "None":
            item["conversion_formula"] = relation

        if used_in and used_in != "N/A" and used_in != "None":
            item["meaning"] = used_in
            item["historical_context"] = used_in
            item["used_in"] = [used_in]

        if ref_url and ref_url != "N/A" and ref_url != "None" and ref_url != "—":
            item["references"] = [ref_url]

        item["tags"] = [
            "mizoram",
            "traditional-units",
            sector_slug,
            clean_slug(unit_name),
            category
        ]

        sheet_units.append(item)
    all_measurements.extend(sheet_units)
    print(f"Sector '{sector_slug}' ({sheet_name}): {len(sheet_units)} units")

print(f"\nTotal Mizoram Measurements: {len(all_measurements)}")

# Check uniqueness of IDs
ids = [m['id'] for m in all_measurements]
assert len(ids) == len(set(ids)), "Duplicate IDs found!"
print("All IDs are unique!")

# Generate TS code
ts_code = '''import { Measurement } from "@/types";

export const MIZORAM_MEASUREMENTS: Measurement[] = ''' + json.dumps(all_measurements, ensure_ascii=False, indent=2) + ';\n'

with open('lib/mizoramData.ts', 'w', encoding='utf-8') as f:
    f.write(ts_code)

print("Successfully written to lib/mizoramData.ts!")
