import openpyxl
import json
import re
import os

path = r'C:\Users\pavan\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\F9C8E35F817C89D982A1A2C62B262666B2E51205\transfers\2026-37\Tamil_Nadu_Traditional_Units (2).xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

sheet_mapping = {
    'Transportation & Distance': ('transportation-distance', 'trans'),
    'Land Measurement': ('land-measurement', 'land'),
    'Livestock & Dairy': ('livestock-dairy', 'dairy'),
    'Household': ('household', 'hh'),
    'Gold & Jewellery': ('gold-jewellery', 'gold'),
    'Agriculture': ('agriculture', 'agri'),
    'Currency & Money': ('currency-money', 'curr'),
    'Storage & Transportation': ('storage-transport', 'storage'),
    'Trade & Commerce': ('trade-commerce', 'trade'),
    'Textile & Handloom': ('textile-handloom', 'textile'),
    'Medicine (Siddha & Ayurveda)': ('medicine', 'med'),
    'Architecture': ('architecture', 'arch'),
}

def clean_slug(text):
    s = re.sub(r'[^a-zA-Z0-9]+', '-', str(text).lower()).strip('-')
    return s if s else "unit"

def map_category(cat_str):
    if not cat_str or cat_str == "N/A": return "other"
    c = str(cat_str).lower()
    if "weight" in c: return "weight"
    if "volume" in c or "capacity" in c: return "volume"
    if "length" in c or "distance" in c: return "length"
    if "area" in c or "land" in c: return "area"
    if "currency" in c or "exchange" in c or "coin" in c or "money" in c: return "currency"
    if "time" in c: return "time"
    if "seed" in c: return "count"
    return "other"

all_measurements = []

for sheet_name, (sector_slug, sector_code) in sheet_mapping.items():
    ws = wb[sheet_name]
    sheet_units = []
    
    for r in range(2, ws.max_row + 1):
        unit_name_val = ws.cell(r, 2).value
        if unit_name_val is None or not str(unit_name_val).strip():
            continue
        
        unit_name = str(unit_name_val).strip()
        tamil_name = str(ws.cell(r, 3).value).strip() if ws.cell(r, 3).value is not None else ""
        english_trans = str(ws.cell(r, 4).value).strip() if ws.cell(r, 4).value is not None else ""
        sanskrit_name = str(ws.cell(r, 5).value).strip() if ws.cell(r, 5).value is not None else ""
        meas_category = str(ws.cell(r, 6).value).strip() if ws.cell(r, 6).value is not None else ""
        order_val = ws.cell(r, 7).value
        relation = str(ws.cell(r, 8).value).strip() if ws.cell(r, 8).value is not None else ""
        modern_si = str(ws.cell(r, 9).value).strip() if ws.cell(r, 9).value is not None else ""
        hist_usage = str(ws.cell(r, 10).value).strip() if ws.cell(r, 10).value is not None else ""
        sector_cell = str(ws.cell(r, 11).value).strip() if ws.cell(r, 11).value is not None else ""
        description = str(ws.cell(r, 12).value).strip() if ws.cell(r, 12).value is not None else ""
        hist_period = str(ws.cell(r, 13).value).strip() if ws.cell(r, 13).value is not None else ""
        region_tn = str(ws.cell(r, 14).value).strip() if ws.cell(r, 14).value is not None else "Tamil Nadu"
        source_ref = str(ws.cell(r, 15).value).strip() if ws.cell(r, 15).value is not None else ""

        row_idx = len(sheet_units) + 1
        item_id = f"tn-{sector_code}-{row_idx}"
        slug = f"tn-{sector_code}-{clean_slug(unit_name)}-{row_idx}"
        category = map_category(meas_category)

        item = {
            "id": item_id,
            "slug": slug,
            "name_english": unit_name,
            "category": category,
            "sector": sector_slug,
            "origin": region_tn if region_tn and region_tn != "None" else "Tamil Nadu",
            "states": ["Tamil Nadu"],
            "created_at": "2024-01-01"
        }

        # Sanskrit Name
        if sanskrit_name and sanskrit_name != "None" and sanskrit_name != "N/A":
            item["name_sanskrit"] = sanskrit_name

        # Tamil Name
        if tamil_name and tamil_name != "—" and tamil_name != "None":
            item["local_names"] = [tamil_name]

        # English Transliteration / Hindi Name mapped to standard table
        if english_trans and english_trans != "—" and english_trans != "None":
            item["name_hindi"] = english_trans
        elif tamil_name and tamil_name != "—" and tamil_name != "None":
            item["name_hindi"] = tamil_name

        # Type / Category
        if meas_category and meas_category != "None":
            item["measurement_type"] = meas_category

        # Approx. Modern Equivalent
        if modern_si and modern_si != "None" and modern_si != "—":
            item["modern_equivalent"] = modern_si

        # Relation / Hierarchy
        if relation and relation != "None" and relation != "—":
            item["conversion_formula"] = relation

        # Meaning & Historical Usage
        if description and description != "None":
            item["meaning"] = description
        
        if hist_usage and hist_usage != "None" and hist_usage != "—":
            item["historical_context"] = hist_usage
            item["used_in"] = [hist_usage]
        elif description:
            item["used_in"] = [description]

        # Historical Period
        if hist_period and hist_period != "None" and hist_period != "—":
            item["historical_period"] = hist_period

        # Region Applicable
        if region_tn and region_tn != "None" and region_tn != "—":
            item["region_applicable"] = region_tn

        # References
        if source_ref and source_ref != "None" and source_ref != "—":
            refs = [r.strip() for r in source_ref.split(';') if r.strip()]
            item["references"] = refs

        item["tags"] = [
            "tamil-nadu",
            "traditional-units",
            sector_slug,
            clean_slug(unit_name),
            category
        ]

        sheet_units.append(item)
    
    all_measurements.extend(sheet_units)
    print(f"Sector '{sector_slug}' ({sheet_name}): {len(sheet_units)} units")

print(f"\nTotal Tamil Nadu Measurements: {len(all_measurements)}")

# Check uniqueness of IDs and Slugs
ids = [m['id'] for m in all_measurements]
assert len(ids) == len(set(ids)), "Duplicate IDs found!"
slugs = [m['slug'] for m in all_measurements]
assert len(slugs) == len(set(slugs)), "Duplicate Slugs found!"
print("All IDs and Slugs are completely unique!")

# Generate TS code
ts_code = '''import { Measurement } from "@/types";

export const TAMILNADU_MEASUREMENTS: Measurement[] = ''' + json.dumps(all_measurements, ensure_ascii=False, indent=2) + ';\n'

with open('lib/tamilnaduData.ts', 'w', encoding='utf-8') as f:
    f.write(ts_code)

print("Successfully written to lib/tamilnaduData.ts!")
