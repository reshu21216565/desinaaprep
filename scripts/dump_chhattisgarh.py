import openpyxl
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')
path = r'C:\Users\pavan\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\F9C8E35F817C89D982A1A2C62B262666B2E51205\transfers\2026-37\chattisgarh.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

for s in wb.sheetnames:
    ws = wb[s]
    headers = [str(ws.cell(1, c).value).strip() if ws.cell(1, c).value is not None else '' for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 2).value is not None and str(ws.cell(r, 2).value).strip():
            row_dict = {}
            for c, h in enumerate(headers, 1):
                v = ws.cell(r, c).value
                row_dict[h] = str(v).strip() if v is not None else ''
            rows.append(row_dict)
    print(f"\n==================== SHEET: '{s}' (Rows: {len(rows)}) ====================")
    print("Headers:", headers)
    if rows:
        print("First row keys & values:")
        for k, v in rows[0].items():
            print(f"  {k}: {v}")
