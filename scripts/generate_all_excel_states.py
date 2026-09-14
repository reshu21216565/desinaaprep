import openpyxl
import json
import re
import os

sheet_to_sector = {
    'Trade & Commerce': ('trade-commerce', 'trade'),
    'Textile & Handloom': ('textile-handloom', 'textile'),
    'Medicine (Ayurveda)': ('medicine', 'med'),
    'Construction & Architecture': ('architecture', 'arch'),
    'Transportation & Distance': ('transportation-distance', 'trans'),
    'Land Measurement': ('land-measurement', 'land'),
    'Livestock & Dairy': ('livestock-dairy', 'dairy'),
    'Household & Daily Life': ('household', 'hh'),
    'Gold & Jewellery': ('gold-jewellery', 'gold'),
    'Seed & Crop (Agriculture)': ('agriculture', 'agri'),
    'Currency & Money': ('currency-money', 'curr'),
    'Storage & Transportation': ('storage-transport', 'storage'),
    'Religious & Cultural Sectors': ('religious-cultural', 'relig')
}

def clean_slug(text):
    s = re.sub(r'[^a-zA-Z0-9]+', '-', text.lower()).strip('-')
    return s

def map_category(type_meas):
    if not type_meas:
        return "other"
    t = type_meas.lower()
    if "weight" in t:
        return "weight"
    if "length" in t or "distance" in t:
        return "length"
    if "volume" in t:
        return "volume"
    if "area" in t:
        return "area"
    if "currency" in t:
        return "currency"
    if "time" in t:
        return "time"
    return "other"

def process_state(file_path, state_name, state_prefix, var_name, output_file):
    wb = openpyxl.load_workbook(file_path)
    entries = []
    
    for sheetname in wb.sheetnames:
        sector_info = sheet_to_sector.get(sheetname)
        if not sector_info:
            print(f"Skipping unknown sheet {sheetname} in {state_name}")
            continue
        sector_slug, sector_code = sector_info
        ws = wb[sheetname]
        
        sheet_entries = []
        for r in range(3, ws.max_row + 1):
            unit_name = ws.cell(r, 2).value
            if unit_name is None:
                continue
            unit_name = str(unit_name).strip()
            if not unit_name:
                continue
            sanskrit = str(ws.cell(r, 3).value).strip() if ws.cell(r, 3).value is not None else ""
            local_name = str(ws.cell(r, 4).value).strip() if ws.cell(r, 4).value is not None else ""
            hindi_name = str(ws.cell(r, 5).value).strip() if ws.cell(r, 5).value is not None else ""
            type_meas = str(ws.cell(r, 6).value).strip() if ws.cell(r, 6).value is not None else ""
            approx_eq = str(ws.cell(r, 7).value).strip() if ws.cell(r, 7).value is not None else ""
            relation = str(ws.cell(r, 8).value).strip() if ws.cell(r, 8).value is not None else ""
            used_in = str(ws.cell(r, 9).value).strip() if ws.cell(r, 9).value is not None else ""
            ref = str(ws.cell(r, 10).value).strip() if ws.cell(r, 10).value is not None else ""
            
            row_num = len(sheet_entries) + 1
            item_id = f"{state_prefix}-{sector_code}-{row_num}"
            slug = f"{clean_slug(unit_name)}-{sector_code}-{state_prefix}"
            
            category = map_category(type_meas)
            
            entry = {
                "id": item_id,
                "slug": slug,
                "name_english": unit_name,
                "category": category,
                "sector": sector_slug,
                "origin": state_name,
                "states": [state_name],
                "local_names": [local_name] if local_name and local_name != "N/A" else [],
                "used_in": [used_in] if used_in else [],
                "references": [ref] if ref and ref != "N/A" else [],
                "tags": [
                    clean_slug(state_name),
                    "traditional-units",
                    sector_slug,
                    clean_slug(unit_name),
                    category
                ],
                "created_at": "2024-01-01"
            }
            if sanskrit and sanskrit != "N/A":
                entry["name_sanskrit"] = sanskrit
            if hindi_name and hindi_name != "N/A":
                entry["name_hindi"] = hindi_name
            if used_in:
                entry["meaning"] = used_in
            if type_meas:
                entry["measurement_type"] = type_meas
            if approx_eq:
                entry["modern_equivalent"] = approx_eq
            if relation:
                entry["conversion_formula"] = relation
            sheet_entries.append(entry)
        entries.extend(sheet_entries)
        print(f"[{state_name}] Sheet '{sheetname}' -> {len(sheet_entries)} entries")
        
    print(f"[{state_name}] Total entries: {len(entries)}")
    ts_code = f'import {{ Measurement }} from "@/types";\n\nexport const {var_name}: Measurement[] = ' + json.dumps(entries, ensure_ascii=False, indent=2) + ';\n'
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write(ts_code)
    print(f"Written to {output_file}\n")

# Process Himachal Pradesh
process_state(
    r'c:\Users\pavan\Downloads\IKS_Traditional_Measurements_HimachalPradesh.xlsx',
    'Himachal Pradesh',
    'hp',
    'HIMACHAL_PRADESH_MEASUREMENTS',
    'lib/himachalPradeshData.ts'
)

# Process Sikkim
process_state(
    r'c:\Users\pavan\Downloads\IKS_Traditional_Measurements_Sikkim.xlsx',
    'Sikkim',
    'sk',
    'SIKKIM_MEASUREMENTS',
    'lib/sikkimData.ts'
)

# Process Uttarakhand
process_state(
    r'c:\Users\pavan\Downloads\IKS_Traditional_Measurements_Uttarakhand.xlsx',
    'Uttarakhand',
    'ut',
    'UTTARAKHAND_MEASUREMENTS',
    'lib/uttarakhandData.ts'
)
