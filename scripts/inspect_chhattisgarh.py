import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
path = r'C:\Users\pavan\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\F9C8E35F817C89D982A1A2C62B262666B2E51205\transfers\2026-37\chattisgarh.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

print(f"Total Sheets in chattisgarh.xlsx: {len(wb.sheetnames)}")
print("Sheet names:", wb.sheetnames)

total_units = 0
for idx, s in enumerate(wb.sheetnames):
    ws = wb[s]
    # Check rows where Traditional Unit Name (col 2) is present
    data_rows = []
    # Check headers row (could be row 1)
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if v is not None and str(v).strip():
            row_data = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
            data_rows.append(row_data)
    total_units += len(data_rows)
    sector_in_row = data_rows[0][9] if (data_rows and len(data_rows[0]) > 9) else "None"
    print(f"\n[{idx+1}] Sheet: '{s}' | Rows: {len(data_rows)} | Headers count: {len(headers)}")
    print(f"  Headers: {headers[:15]}")
    if data_rows:
        print(f"  Row 1: {data_rows[0][:5]}")

print(f"\nTotal units across all sheets in chattisgarh.xlsx: {total_units}")
